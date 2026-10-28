from __future__ import annotations

import re
from pathlib import Path
from typing import Iterable

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[2]
SOURCE = ROOT / "outputs" / "武汉光谷无人机配送_PPO开发指南.md"
OUTPUT = Path(__file__).with_name("无人机项目设计开发指南_v3.0.docx")
IMAGE_DIR = ROOT / "04_脚本图片输出" / "论文插图"

NAVY = "0B2545"
BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
MUTED = "667085"
LIGHT_BLUE = "E8EEF5"
LIGHT_GRAY = "F4F6F9"
BORDER = "B8C4D2"
WHITE = "FFFFFF"
BLACK = "1F2937"

CONTENT_DXA = 9360
TABLE_INDENT_DXA = 120
CELL_TOP_BOTTOM_DXA = 80
CELL_SIDE_DXA = 120


def rgb(value: str) -> RGBColor:
    return RGBColor.from_string(value)


def set_run_font(
    run,
    *,
    latin: str = "Calibri",
    east_asia: str = "Microsoft YaHei",
    size: float | None = None,
    color: str | None = None,
    bold: bool | None = None,
    italic: bool | None = None,
):
    run.font.name = latin
    rpr = run._element.get_or_add_rPr()
    fonts = rpr.rFonts
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.insert(0, fonts)
    fonts.set(qn("w:ascii"), latin)
    fonts.set(qn("w:hAnsi"), latin)
    fonts.set(qn("w:eastAsia"), east_asia)
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = rgb(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def set_cell_shading(cell, fill: str):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.find(qn("w:tcMar"))
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    values = {
        "top": CELL_TOP_BOTTOM_DXA,
        "bottom": CELL_TOP_BOTTOM_DXA,
        "start": CELL_SIDE_DXA,
        "end": CELL_SIDE_DXA,
    }
    for side, value in values.items():
        node = tc_mar.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_table_borders(table, color: str = BORDER, size: int = 6):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = qn(f"w:{edge}")
        node = borders.find(tag)
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), str(size))
        node.set(qn("w:space"), "0")
        node.set(qn("w:color"), color)


def apply_table_geometry(table, widths_dxa: list[int]):
    total = sum(widths_dxa)
    tbl_pr = table._tbl.tblPr

    tbl_layout = tbl_pr.find(qn("w:tblLayout"))
    if tbl_layout is None:
        tbl_layout = OxmlElement("w:tblLayout")
        tbl_pr.append(tbl_layout)
    tbl_layout.set(qn("w:type"), "fixed")

    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(total))
    tbl_w.set(qn("w:type"), "dxa")

    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(TABLE_INDENT_DXA))
    tbl_ind.set(qn("w:type"), "dxa")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(widths_dxa[idx]))
            tc_w.set(qn("w:type"), "dxa")
            cell.width = Inches(widths_dxa[idx] / 1440)
            set_cell_margins(cell)


def add_field(paragraph, instruction: str):
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, text, end])
    set_run_font(run, size=9, color=MUTED)


def set_paragraph_shading(paragraph, fill: str, border_color: str | None = None):
    p_pr = paragraph._p.get_or_add_pPr()
    shd = p_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        p_pr.append(shd)
    shd.set(qn("w:fill"), fill)
    if border_color:
        borders = p_pr.find(qn("w:pBdr"))
        if borders is None:
            borders = OxmlElement("w:pBdr")
            p_pr.append(borders)
        left = OxmlElement("w:left")
        left.set(qn("w:val"), "single")
        left.set(qn("w:sz"), "18")
        left.set(qn("w:space"), "6")
        left.set(qn("w:color"), border_color)
        borders.append(left)


def style_document(doc: Document):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)
    section.different_first_page_header_footer = True

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.font.color.rgb = rgb(BLACK)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25
    normal.paragraph_format.keep_together = False
    normal.paragraph_format.widow_control = True

    heading_tokens = {
        "Heading 1": (16, BLUE, 18, 10),
        "Heading 2": (13, BLUE, 14, 7),
        "Heading 3": (12, DARK_BLUE, 10, 5),
    }
    for name, (size, color, before, after) in heading_tokens.items():
        style = styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = rgb(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.line_spacing = 1.0
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.keep_together = True

    for name in ("List Bullet", "List Number"):
        style = styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(11)
        style.paragraph_format.left_indent = Inches(0.375)
        style.paragraph_format.first_line_indent = Inches(-0.188)
        style.paragraph_format.space_before = Pt(0)
        style.paragraph_format.space_after = Pt(4)
        style.paragraph_format.line_spacing = 1.25

    if "Code Block" not in styles:
        code_style = styles.add_style("Code Block", WD_STYLE_TYPE.PARAGRAPH)
    else:
        code_style = styles["Code Block"]
    code_style.font.name = "Consolas"
    code_style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    code_style.font.size = Pt(8.5)
    code_style.font.color.rgb = rgb("243047")
    code_style.paragraph_format.left_indent = Inches(0.14)
    code_style.paragraph_format.right_indent = Inches(0.08)
    code_style.paragraph_format.space_before = Pt(0)
    code_style.paragraph_format.space_after = Pt(0)
    code_style.paragraph_format.line_spacing = 1.05

    if "Caption Custom" not in styles:
        caption = styles.add_style("Caption Custom", WD_STYLE_TYPE.PARAGRAPH)
    else:
        caption = styles["Caption Custom"]
    caption.font.name = "Calibri"
    caption._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    caption.font.size = Pt(9)
    caption.font.italic = True
    caption.font.color.rgb = rgb(MUTED)
    caption.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.paragraph_format.space_before = Pt(4)
    caption.paragraph_format.space_after = Pt(9)
    caption.paragraph_format.keep_with_next = False


def add_header_footer(section):
    header = section.header
    p = header.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_after = Pt(0)
    run = p.add_run("无人机项目设计开发指南")
    set_run_font(run, size=9, color=MUTED, bold=True)

    footer = section.footer
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    fp.paragraph_format.space_before = Pt(0)
    fp.paragraph_format.space_after = Pt(0)
    run = fp.add_run("第 ")
    set_run_font(run, size=9, color=MUTED)
    add_field(fp, "PAGE")
    run = fp.add_run(" 页")
    set_run_font(run, size=9, color=MUTED)


def add_cover(doc: Document):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(0)
    p.add_run().add_break()
    p.add_run().add_break()
    p.add_run().add_break()

    kicker = doc.add_paragraph()
    kicker.alignment = WD_ALIGN_PARAGRAPH.CENTER
    kicker.paragraph_format.space_after = Pt(18)
    run = kicker.add_run("技术参考手册")
    set_run_font(run, size=11, color=BLUE, bold=True)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_before = Pt(0)
    title.paragraph_format.space_after = Pt(12)
    title.paragraph_format.keep_with_next = True
    run = title.add_run("无人机项目设计\n开发指南")
    set_run_font(run, size=30, color=NAVY, bold=True)

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(34)
    run = subtitle.add_run("基于深度强化学习的多无人机协同物流配送路径规划")
    set_run_font(run, size=14, color=DARK_BLUE)

    meta = doc.add_table(rows=4, cols=2)
    meta.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta.autofit = False
    rows = [
        ("文档版本", "v3.0"),
        ("编制日期", "2026-07-05"),
        ("适用对象", "算法开发、仿真训练、实验复现人员"),
        ("内容依据", "当前源码、四区场景、测试、开悟指南与 Gorge Walk v2"),
    ]
    for idx, (label, value) in enumerate(rows):
        meta.cell(idx, 0).text = label
        meta.cell(idx, 1).text = value
        set_cell_shading(meta.cell(idx, 0), LIGHT_BLUE)
        for cell in meta.rows[idx].cells:
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_before = Pt(2)
                paragraph.paragraph_format.space_after = Pt(2)
                for run in paragraph.runs:
                    set_run_font(
                        run,
                        size=10,
                        color=NAVY if cell == meta.cell(idx, 0) else BLACK,
                        bold=cell == meta.cell(idx, 0),
                    )
    set_repeat_table_header(meta.rows[0])
    apply_table_geometry(meta, [1800, 5400])
    set_table_borders(meta, color="CCD6E0", size=5)

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(32)
    p.paragraph_format.space_after = Pt(0)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("算法与仿真研究用途 · 不涉及真实飞控与硬件部署")
    set_run_font(r, size=9.5, color=MUTED, italic=True)

    doc.add_page_break()


def add_toc(doc: Document):
    p = doc.add_paragraph("使用导航", style="Heading 1")
    p.paragraph_format.space_before = Pt(0)
    lead = doc.add_paragraph(
        "本指南先给出项目边界和可运行主链路，再展开环境、算法、训练、实验与扩展开发。"
    )
    lead.paragraph_format.space_after = Pt(10)
    sections = [
        "第一部分. 场景、环境、任务、奖励、PPO 与实验",
        "0–5. 项目结论、场景资产与环境数据",
        "6–9. 多目标配送、返航状态机与 PPO",
        "10–16. 课程、训练、评估、测试与实施",
        "第二部分. 开悟式强化学习开发框架",
        "17–20. 框架、环境、场景包与规划器 API",
        "21–24. Agent、数据协议、算法与工作流",
        "25–31. 评估、分布式、监控、测试与安全",
        "第三部分. 文件导航、API 速查与发布检查表",
    ]
    table = doc.add_table(rows=len(sections), cols=2)
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for idx, item in enumerate(sections):
        label, title = item.split(". ", 1)
        table.cell(idx, 0).text = label
        table.cell(idx, 1).text = title
        if idx % 2 == 0:
            set_cell_shading(table.cell(idx, 0), LIGHT_GRAY)
            set_cell_shading(table.cell(idx, 1), LIGHT_GRAY)
        for cell in table.rows[idx].cells:
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            for para in cell.paragraphs:
                para.paragraph_format.space_before = Pt(1)
                para.paragraph_format.space_after = Pt(1)
                for run in para.runs:
                    set_run_font(
                        run,
                        size=9.5,
                        color=BLUE if cell == table.cell(idx, 0) else BLACK,
                        bold=cell == table.cell(idx, 0),
                    )
    set_repeat_table_header(table.rows[0])
    apply_table_geometry(table, [1050, 7260])
    set_table_borders(table, color="D8E0E8", size=4)
    doc.add_page_break()


INLINE_RE = re.compile(r"(\*\*.+?\*\*|`.+?`)")


def add_inline_runs(paragraph, text: str, *, default_size: float = 11):
    pos = 0
    for match in INLINE_RE.finditer(text):
        if match.start() > pos:
            run = paragraph.add_run(text[pos : match.start()])
            set_run_font(run, size=default_size, color=BLACK)
        token = match.group(0)
        if token.startswith("**"):
            run = paragraph.add_run(token[2:-2])
            set_run_font(run, size=default_size, color=NAVY, bold=True)
        else:
            run = paragraph.add_run(token[1:-1])
            set_run_font(
                run,
                latin="Consolas",
                east_asia="Microsoft YaHei",
                size=max(8.5, default_size - 1),
                color=DARK_BLUE,
            )
        pos = match.end()
    if pos < len(text):
        run = paragraph.add_run(text[pos:])
        set_run_font(run, size=default_size, color=BLACK)


def add_code_block(doc: Document, lines: Iterable[str]):
    lines = list(lines)
    for idx, line in enumerate(lines or [""]):
        p = doc.add_paragraph(style="Code Block")
        if idx == 0:
            p.paragraph_format.space_before = Pt(5)
        if idx == len(lines) - 1:
            p.paragraph_format.space_after = Pt(7)
        set_paragraph_shading(p, LIGHT_GRAY)
        run = p.add_run(line if line else " ")
        set_run_font(
            run,
            latin="Consolas",
            east_asia="Microsoft YaHei",
            size=8.5,
            color="243047",
        )


def parse_table_row(line: str) -> list[str]:
    return [cell.strip() for cell in line.strip().strip("|").split("|")]


def choose_widths(rows: list[list[str]]) -> list[int]:
    cols = len(rows[0])
    scores = []
    for idx in range(cols):
        max_len = max(len(row[idx]) if idx < len(row) else 0 for row in rows)
        avg_len = sum(len(row[idx]) if idx < len(row) else 0 for row in rows) / len(rows)
        scores.append(max(5.0, min(28.0, 0.65 * max_len + 0.35 * avg_len)))
    total_score = sum(scores)
    widths = [int(CONTENT_DXA * score / total_score) for score in scores]
    min_width = 900 if cols <= 4 else 650
    for idx, width in enumerate(widths):
        widths[idx] = max(min_width, width)
    scale = CONTENT_DXA / sum(widths)
    widths = [int(width * scale) for width in widths]
    widths[-1] += CONTENT_DXA - sum(widths)
    return widths


def add_markdown_table(doc: Document, rows: list[list[str]]):
    cols = len(rows[0])
    rows = [row + [""] * (cols - len(row)) for row in rows]
    table = doc.add_table(rows=len(rows), cols=cols)
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    for r_idx, row in enumerate(rows):
        for c_idx, text in enumerate(row):
            cell = table.cell(r_idx, c_idx)
            cell.text = ""
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            if r_idx == 0:
                set_cell_shading(cell, LIGHT_BLUE)
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_before = Pt(2)
            paragraph.paragraph_format.space_after = Pt(2)
            paragraph.paragraph_format.line_spacing = 1.08
            add_inline_runs(paragraph, text, default_size=8.8 if cols >= 4 else 9.3)
            if r_idx == 0:
                for run in paragraph.runs:
                    run.bold = True
                    run.font.color.rgb = rgb(NAVY)
    set_repeat_table_header(table.rows[0])
    apply_table_geometry(table, choose_widths(rows))
    set_table_borders(table)
    after = doc.add_paragraph()
    after.paragraph_format.space_before = Pt(0)
    after.paragraph_format.space_after = Pt(2)


def add_figure_path(doc: Document, path: Path, caption: str, width: float):
    if not path.exists():
        return
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.keep_with_next = True
    run = p.add_run()
    shape = run.add_picture(str(path), width=Inches(width))
    shape._inline.docPr.set("descr", caption)
    shape._inline.docPr.set("title", path.name)
    cp = doc.add_paragraph(caption, style="Caption Custom")
    cp.paragraph_format.keep_with_next = False


def add_figure(doc: Document, filename: str, caption: str, width: float):
    add_figure_path(doc, IMAGE_DIR / filename, caption, width)


def add_heading(doc: Document, level: int, text: str):
    if text in {
        "第一部分 场景与智能体开发指南",
        "第二部分 开悟式强化学习开发框架",
        "第三部分 附录",
    }:
        if len(doc.paragraphs) > 2:
            doc.add_page_break()
    p = doc.add_paragraph(style=f"Heading {min(level, 3)}")
    add_inline_runs(p, text, default_size={1: 16, 2: 13, 3: 12}.get(level, 11))
    for run in p.runs:
        run.bold = True
        run.font.color.rgb = rgb(BLUE if level < 3 else DARK_BLUE)

    if text.startswith("1.4 "):
        add_figure(
            doc,
            "algorithm_architecture.png",
            "图 1  PPO + CTDE 目标架构示意（CTDE 当前仍为独立原型）",
            6.15,
        )
    elif text.startswith("6.1 "):
        add_figure(
            doc,
            "environment_3d_view.png",
            "图 2  100 × 100 × 50 m 三维空域示意",
            5.75,
        )
    elif text.startswith("6.2 "):
        add_figure(
            doc,
            "urban_environment_3d_enhanced.png",
            "图 3  2 km 城市复杂空域与示意配送路径",
            5.15,
        )


def markdown_to_doc(doc: Document, source: str):
    lines = source.splitlines()
    idx = 0
    in_code = False
    code_lines: list[str] = []
    skip_front = True

    while idx < len(lines):
        line = lines[idx].rstrip()

        if skip_front:
            if line.startswith("# 第一部分"):
                skip_front = False
            else:
                idx += 1
                continue

        if line.startswith("```"):
            if in_code:
                add_code_block(doc, code_lines)
                code_lines = []
                in_code = False
            else:
                in_code = True
            idx += 1
            continue
        if in_code:
            code_lines.append(line)
            idx += 1
            continue

        if not line.strip() or line.strip() == "---":
            idx += 1
            continue

        table_start = line.lstrip().startswith("|")
        if table_start and idx + 1 < len(lines):
            delimiter = lines[idx + 1].strip()
            if re.match(r"^\|?[\s:|-]+\|?$", delimiter):
                rows = [parse_table_row(line)]
                idx += 2
                while idx < len(lines) and lines[idx].lstrip().startswith("|"):
                    rows.append(parse_table_row(lines[idx]))
                    idx += 1
                add_markdown_table(doc, rows)
                continue

        image = re.match(r"^!\[(.+?)\]\((.+?)\)$", line)
        if image:
            image_path = (SOURCE.parent / image.group(2)).resolve()
            add_figure_path(
                doc, image_path, image.group(1), width=6.1
            )
            idx += 1
            continue

        heading = re.match(r"^(#{1,4})\s+(.+)$", line)
        if heading:
            hashes = len(heading.group(1))
            level = 1 if hashes <= 2 else hashes - 1
            add_heading(doc, level, heading.group(2))
            idx += 1
            continue

        if line.startswith("> "):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.18)
            p.paragraph_format.right_indent = Inches(0.08)
            p.paragraph_format.space_before = Pt(5)
            p.paragraph_format.space_after = Pt(7)
            set_paragraph_shading(p, "EDF4FB", BLUE)
            add_inline_runs(p, line[2:])
            idx += 1
            continue

        bullet = re.match(r"^\s*-\s+(.+)$", line)
        if bullet:
            p = doc.add_paragraph(style="List Bullet")
            add_inline_runs(p, bullet.group(1))
            idx += 1
            continue

        numbered = re.match(r"^\s*\d+\.\s+(.+)$", line)
        if numbered:
            p = doc.add_paragraph(style="List Number")
            add_inline_runs(p, numbered.group(1))
            idx += 1
            continue

        p = doc.add_paragraph()
        add_inline_runs(p, line)
        idx += 1


def add_document_properties(doc: Document):
    props = doc.core_properties
    props.title = "无人机项目设计开发指南"
    props.subject = "开悟范式下的多目标配送、安全返航与 PPO 开发"
    props.author = "无人机项目组"
    props.keywords = "无人机, PPO, CTDE, 强化学习, 路径规划, 仿真"
    props.comments = "依据 2026-07-05 工作区源码、场景、测试与参考指南编制"


def build():
    doc = Document()
    style_document(doc)
    add_document_properties(doc)
    add_header_footer(doc.sections[0])
    add_cover(doc)
    add_toc(doc)
    markdown_to_doc(doc, SOURCE.read_text(encoding="utf-8"))

    for paragraph in doc.paragraphs:
        paragraph.paragraph_format.widow_control = True

    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
